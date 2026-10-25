"""Validate two blinded ConsumerBR pilot workbooks; never infer missing labels."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PRIVATE = ROOT / "outputs/tp3-consumerbr"
SOURCE = PRIVATE / "pilot_development.csv"
MANIFEST = ROOT / "reports/tp3_data/consumerbr_pilot_manifest.json"
TAXONOMY = ROOT / "data/processed/bitext_retail_intents_report.json"
FIRST_ROW = 7


def load_source() -> dict[str, str]:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    if hashlib.sha256(SOURCE.read_bytes()).hexdigest() != manifest["development_file_sha256"]:
        raise ValueError("CSV privado não corresponde ao manifesto fixado.")
    with SOURCE.open(encoding="utf-8", newline="") as source:
        rows = list(csv.DictReader(source))
    if len(rows) != 240 or len({row["sample_id"] for row in rows}) != 240:
        raise ValueError("Fila de desenvolvimento deve conter 240 IDs únicos.")
    return {row["sample_id"]: row["complaint_text"] for row in rows}


def allowed_labels() -> set[str]:
    report = json.loads(TAXONOMY.read_text(encoding="utf-8"))
    labels = set(report["intent_counts"]) | {"fora_de_escopo"}
    if len(labels) != 47:
        raise ValueError("Taxonomia deve ter 46 intenções Bitext e fora_de_escopo.")
    return labels


def normalize_cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def validate_record(record: dict, labels: set[str], *, require_complete: bool) -> list[str]:
    """Return actionable errors; uncertain cases may legitimately lack an intent."""
    errors: list[str] = []
    intent = record["intent"]
    uncertain = record["uncertain"]
    if intent and intent not in labels:
        errors.append("intenção fora da taxonomia")
    if uncertain and uncertain not in {"sim", "não"}:
        errors.append("incerto deve ser sim ou não")
    if require_complete:
        if not uncertain:
            errors.append("incerto vazio")
        if uncertain == "não" and not intent:
            errors.append("intenção vazia sem marca de incerteza")
        if (uncertain == "sim" or intent == "fora_de_escopo") and not record["notes"]:
            errors.append("justificativa vazia")
        if not record["reviewer"]:
            errors.append("revisor vazio")
    return errors


def read_workbook(path: Path, source_text: dict[str, str], labels: set[str], *, require_complete: bool) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    if "Anotação" not in workbook.sheetnames:
        raise ValueError(f"{path.name}: aba Anotação ausente.")
    sheet = workbook["Anotação"]
    expected_header = ["ID", "Texto da reclamação", "Intenção", "Incerto?", "Justificativa", "Revisor", "Situação"]
    header = [normalize_cell(cell.value) for cell in sheet["A6:G6"][0]]
    if header != expected_header:
        raise ValueError(f"{path.name}: cabeçalho alterado.")
    records = []
    seen = set()
    problems = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=FIRST_ROW, max_row=FIRST_ROW + 239, max_col=7, values_only=True), start=FIRST_ROW):
        sample_id, text, intent, uncertain, notes, reviewer, _status = (normalize_cell(value) for value in values)
        if sample_id not in source_text:
            problems.append(f"linha {row_number}: ID desconhecido")
            continue
        if sample_id in seen:
            problems.append(f"linha {row_number}: ID duplicado")
        seen.add(sample_id)
        if text != source_text[sample_id]:
            problems.append(f"linha {row_number}: texto de origem alterado")
        record = {"sample_id": sample_id, "intent": intent, "uncertain": uncertain, "notes": notes, "reviewer": reviewer}
        problems.extend(f"linha {row_number}: {issue}" for issue in validate_record(record, labels, require_complete=require_complete))
        records.append(record)
    if seen != set(source_text):
        problems.append(f"faltam {len(set(source_text) - seen)} IDs da fila")
    if problems:
        raise ValueError(f"{path.name}: " + "; ".join(problems[:10]) + (f"; ... ({len(problems)} problemas)" if len(problems) > 10 else ""))
    workbook.close()
    return records


def compare(a: list[dict], b: list[dict]) -> tuple[list[dict], dict]:
    lookup_b = {row["sample_id"]: row for row in b}
    decisions = []
    disagreements = 0
    uncertain = 0
    for left in a:
        right = lookup_b[left["sample_id"]]
        is_uncertain = left["uncertain"] == "sim" or right["uncertain"] == "sim"
        is_disagreement = left["intent"] != right["intent"]
        uncertain += is_uncertain
        disagreements += is_disagreement
        decisions.append({
            "sample_id": left["sample_id"],
            "intent_a": left["intent"], "uncertain_a": left["uncertain"],
            "notes_a": left["notes"], "reviewer_a": left["reviewer"],
            "intent_b": right["intent"], "uncertain_b": right["uncertain"],
            "notes_b": right["notes"], "reviewer_b": right["reviewer"],
            "needs_adjudication": is_uncertain or is_disagreement,
        })
    return decisions, {
        "rows": len(decisions),
        "different_intents": disagreements,
        "any_uncertain": uncertain,
        "needs_adjudication": sum(row["needs_adjudication"] for row in decisions),
        "human_ground_truth_complete": False,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--annotator-a", type=Path, default=PRIVATE / "piloto_anotador_A.xlsx")
    parser.add_argument("--annotator-b", type=Path, default=PRIVATE / "piloto_anotador_B.xlsx")
    parser.add_argument("--check-template", action="store_true", help="Conferir estrutura sem exigir respostas humanas")
    args = parser.parse_args()
    source_text = load_source()
    labels = allowed_labels()
    a = read_workbook(args.annotator_a, source_text, labels, require_complete=not args.check_template)
    b = read_workbook(args.annotator_b, source_text, labels, require_complete=not args.check_template)
    decisions, summary = compare(a, b)
    if args.check_template:
        if any(row["intent_a"] or row["intent_b"] or row["notes_a"] or row["notes_b"] for row in decisions):
            raise ValueError("O template contém respostas e não está em branco.")
        print("Templates A e B válidos: 240 IDs, textos intactos, respostas vazias.")
        return
    out = PRIVATE / "pilot_annotation_comparison.csv"
    with out.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=list(decisions[0]))
        writer.writeheader()
        writer.writerows(decisions)
    (PRIVATE / "pilot_annotation_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))
    print(f"Decisões humanas guardadas apenas no diretório privado: {out}")


if __name__ == "__main__":
    main()
