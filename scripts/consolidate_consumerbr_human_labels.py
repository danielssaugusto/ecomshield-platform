"""Consolidate clear agreements and adjudicated ConsumerBR labels privately."""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PRIVATE = ROOT / "outputs/tp3-consumerbr"
TAXONOMY = ROOT / "data/processed/bitext_retail_intents_report.json"


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def main() -> None:
    labels = set(json.loads(TAXONOMY.read_text(encoding="utf-8"))["intent_counts"]) | {"fora_de_escopo"}
    with (PRIVATE / "pilot_development.csv").open(encoding="utf-8", newline="") as source_file:
        source = {row["sample_id"]: row for row in csv.DictReader(source_file)}
    with (PRIVATE / "pilot_annotation_comparison.csv").open(encoding="utf-8", newline="") as comparison_file:
        comparison = {row["sample_id"]: row for row in csv.DictReader(comparison_file)}

    adjudication_path = PRIVATE / "consumerbr_adjudication_nathalia_artigas.xlsx"
    workbook = load_workbook(adjudication_path, read_only=True, data_only=False)
    sheet = workbook["Adjudicação"]
    adjudicated: dict[str, dict[str, str]] = {}
    for row_number, row in enumerate(sheet.iter_rows(min_row=7, max_row=173, max_col=12, values_only=True), start=7):
        sample_id, text, final_label, reason, adjudicator = clean(row[0]), clean(row[1]), clean(row[8]), clean(row[9]), clean(row[10])
        if sample_id not in source or text != source[sample_id]["complaint_text"]:
            raise ValueError(f"Linha {row_number}: ID ou texto de adjudicação não corresponde à origem.")
        if final_label not in labels or not reason or not adjudicator:
            raise ValueError(f"Linha {row_number}: decisão final, justificativa ou adjudicador inválido.")
        if sample_id in adjudicated:
            raise ValueError(f"Linha {row_number}: ID de adjudicação duplicado.")
        adjudicated[sample_id] = {"label": final_label, "reason": reason, "adjudicator": adjudicator}
    workbook.close()

    rows = []
    agreement_count = 0
    for sample_id, raw in source.items():
        pair = comparison.get(sample_id)
        if pair is None:
            raise ValueError(f"ID {sample_id} ausente da comparação entre anotadores.")
        if sample_id in adjudicated:
            final = adjudicated[sample_id]
            origin = "adjudication"
            justification = final["reason"]
            adjudicator = final["adjudicator"]
        else:
            if pair["needs_adjudication"] != "False" or pair["intent_a"] != pair["intent_b"] or not pair["intent_a"]:
                raise ValueError(f"ID {sample_id} não pode ser aceito como concordância clara.")
            final = {"label": pair["intent_a"]}
            origin = "clear_agreement"
            justification = ""
            adjudicator = ""
            agreement_count += 1
        rows.append({
            "sample_id": sample_id,
            "complaint_text": raw["complaint_text"],
            "final_intent": final["label"],
            "label_origin": origin,
            "luana_intent": pair["intent_a"],
            "luana_uncertain": pair["uncertain_a"],
            "daniel_intent": pair["intent_b"],
            "daniel_uncertain": pair["uncertain_b"],
            "adjudication_justification": justification,
            "adjudicator": adjudicator,
        })
    if len(rows) != 240 or len(adjudicated) != 167 or agreement_count != 73:
        raise ValueError("Contagens de consolidação inesperadas.")
    if {row["final_intent"] for row in rows} - labels:
        raise ValueError("Rótulo final fora da taxonomia.")

    output_csv = PRIVATE / "consumerbr_ptbr_human_labels.csv"
    with output_csv.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    counts = Counter(row["final_intent"] for row in rows)
    report = {
        "development_rows": len(rows),
        "clear_agreement_rows": agreement_count,
        "adjudicated_rows": len(adjudicated),
        "final_intent_coverage": len(counts),
        "reserved_test_rows": 80,
        "reserved_test_used_for_selection": False,
        "final_intent_counts": dict(sorted(counts.items())),
        "limitations": [
            "The 240 development rows are a company-stratified ConsumerBR pilot, not a representative production-support distribution.",
            "The reserved 80 rows remain unlabeled and must not be used for model selection.",
            "The final labels are suitable for pilot training analysis, not a claim of 46-intent coverage unless the reported counts show it.",
        ],
    }
    (PRIVATE / "consumerbr_ptbr_human_labels_summary.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    public_report = ROOT / "reports/tp3_data/consumerbr_human_labels.md"
    public_report.write_text(
        "# Consolidação do piloto ConsumerBR\n\n"
        f"Foram consolidadas **{len(rows)}** reclamações reais anonimizadas para desenvolvimento: "
        f"**{agreement_count}** por concordância clara e **{len(adjudicated)}** por adjudicação humana documentada. "
        f"Os rótulos finais cobrem **{len(counts)}** das 46 intenções Bitext, além de `fora_de_escopo` quando necessário.\n\n"
        "O CSV com texto permanece privado e ignorado pelo Git. A fila de 80 casos reservados continua sem rótulos e não foi usada "
        "na seleção de modelo. Portanto, ainda não há uma nova métrica externa ou evidência de atingir 60%.\n",
        encoding="utf-8",
    )
    print(f"Consolidação concluída: {len(rows)} linhas, {agreement_count} concordâncias, {len(adjudicated)} adjudicações, {len(counts)} rótulos finais.")


if __name__ == "__main__":
    main()
